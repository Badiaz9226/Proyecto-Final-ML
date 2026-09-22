"""Complete the professor's DOCX template with the verified project outputs."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


NAVY = "173A5E"
BLUE = "176B87"
TEAL = "1F8A70"
ORANGE = "D98418"
LIGHT = "EAF3F5"
WHITE = "FFFFFF"


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def fmt(value: Any, digits: int = 4) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "No disponible"
    if isinstance(value, (int,)):
        return f"{value:,}".replace(",", ".")
    if isinstance(value, (float,)):
        return f"{value:.{digits}f}".replace(".", ",")
    return str(value)


def percentage(value: float, digits: int = 2) -> str:
    return f"{100 * value:.{digits}f}%".replace(".", ",")


def clear_cell(cell: Any) -> None:
    cell.text = ""


def set_cell(cell: Any, value: Any, *, bold: bool = False) -> None:
    clear_cell(cell)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(value))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_hyperlink(cell: Any, text: str, url: str) -> None:
    clear_cell(cell)
    paragraph = cell.paragraphs[0]
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def shade(cell: Any, fill: str) -> None:
    tc_properties = cell._tc.get_or_add_tcPr()
    shading = tc_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_properties.append(shading)
    shading.set(qn("w:fill"), fill)


def repeat_header(row: Any) -> None:
    tr_properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_properties.append(header)


def set_table_borders(table: Any, color: str = "A8B5BF", size: str = "4") -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def remove_rows_after(table: Any, keep: int) -> None:
    for row in list(table.rows)[keep:]:
        table._tbl.remove(row._tr)


def set_table_width(table: Any, column_count: int, total_width_cm: float = 18.0) -> None:
    """Force template and appended tables to use the printable page width."""
    table.autofit = False
    table_properties = table._tbl.tblPr
    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:type"), "pct")
    table_width.set(qn("w:w"), "5000")

    width_twips = int(total_width_cm / column_count / 2.54 * 1440)
    for row in table.rows:
        for cell in row.cells[:column_count]:
            cell.width = Cm(total_width_cm / column_count)
            tc_properties = cell._tc.get_or_add_tcPr()
            tc_width = tc_properties.find(qn("w:tcW"))
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                tc_properties.append(tc_width)
            tc_width.set(qn("w:type"), "dxa")
            tc_width.set(qn("w:w"), str(width_twips))


def replace_table(table: Any, headers: list[str], rows: Iterable[Iterable[Any]]) -> None:
    remove_rows_after(table, 1)
    while len(table.columns) < len(headers):
        raise ValueError("The template table has fewer columns than required.")
    for index, header in enumerate(headers):
        set_cell(table.cell(0, index), header, bold=True)
        shade(table.cell(0, index), NAVY)
        for run in table.cell(0, index).paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    repeat_header(table.rows[0])
    for row_index, values in enumerate(rows, start=1):
        row = table.add_row()
        for column_index, value in enumerate(values):
            set_cell(row.cells[column_index], value)
            if row_index % 2 == 0:
                shade(row.cells[column_index], LIGHT)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, len(headers))


def add_table(document: Document, headers: list[str], rows: Iterable[Iterable[Any]]) -> Any:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Normal Table"
    set_table_borders(table)
    replace_table(table, headers, rows)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def set_fixed_table(table: Any, values: list[str], links: dict[int, str] | None = None) -> None:
    links = links or {}
    for index, value in enumerate(values):
        if index in links:
            add_hyperlink(table.cell(index, 1), value, links[index])
        else:
            set_cell(table.cell(index, 1), value)
        shade(table.cell(index, 0), LIGHT)
        set_cell(table.cell(index, 0), table.cell(index, 0).text, bold=True)


def add_page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


def add_heading(document: Document, text: str, level: int = 1) -> Any:
    heading = document.add_heading(text, level=level)
    heading.paragraph_format.keep_with_next = True
    return heading


def add_figure(document: Document, path: Path, caption: str, width: float = 6.2) -> None:
    if not path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.style = document.styles["Caption"]


def get_prompt_rows(notebook_path: Path) -> list[list[str]]:
    if not notebook_path.exists():
        return []
    notebook = read_json(notebook_path)
    rows: list[list[str]] = []
    pattern = re.compile(r"^### Prompt (\d{2}) — (.+)$", re.MULTILINE)
    for cell in notebook.get("cells", []):
        if cell.get("cell_type") != "markdown":
            continue
        text = "".join(cell.get("source", []))
        match = pattern.search(text)
        if match:
            rows.append([match.group(1), match.group(2), "Código y revisión visibles en Colab"])
    return rows


def configure_document(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(4)
    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.color.rgb = RGBColor.from_string(NAVY)
    document.styles["Heading 1"].font.size = Pt(16)
    document.styles["Heading 2"].font.size = Pt(13)
    document.styles["Heading 3"].font.size = Pt(11)
    for section in document.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.55)
        section.right_margin = Cm(1.55)
        header = section.header.paragraphs[0]
        header.text = "Proyecto Integrador · Aprendizaje de Máquinas · CUM"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        add_page_number(section.footer.paragraphs[0])


def build_document(args: argparse.Namespace) -> None:
    root = args.output.parent.parent
    reports = root / "reports"
    figures = root / "figures"
    results = read_json(reports / "resultados_modelado.json")
    quality = read_json(reports / "resumen_calidad.json")
    dictionary = pd.read_csv(reports / "diccionario_datos_original.csv")
    quality_table = pd.read_csv(reports / "calidad_datos.csv")
    associations = pd.read_csv(reports / "asociaciones_variables.csv")
    cv = pd.read_csv(reports / "metricas_validacion_cruzada.csv")
    test = pd.read_csv(reports / "metricas_test.csv")
    predictions = pd.read_csv(reports / "predicciones_datos_no_vistos.csv")

    document = Document(args.template)
    if document.paragraphs and "Nota para diseño" in document.paragraphs[0].text:
        paragraph = document.paragraphs[0]
        paragraph._element.getparent().remove(paragraph._element)
    configure_document(document)

    title = document.paragraphs[0]
    title.text = "Formato de Entrega\nProyecto Integrador"
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.paragraphs[1].insert_paragraph_before(
        "Clasificación de presentaciones del Código Único de Medicamentos como muestra médica\n"
        "Brayam Arboleda Diaz · Saul Dario Gomez · Corte 21/09/2026"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.name = "Aptos"
        run.font.color.rgb = RGBColor.from_string(BLUE)
        run.font.size = Pt(10)

    source_url = quality["fuente"]["url_oficial"]
    general_values = [
        "Clasificación de presentaciones CUM como muestra médica",
        "Brayam Arboleda Diaz y Saul Dario Gomez",
        "Clasificación binaria supervisada",
        "INVIMA — Datos Abiertos Colombia",
        source_url,
        args.colab_url,
        args.github_url,
        args.streamlit_url,
    ]
    set_fixed_table(
        document.tables[0],
        general_values,
        links={4: source_url, 5: args.colab_url, 6: args.github_url, 7: args.streamlit_url},
    )

    business = [
        "El INVIMA publica información administrativa de los códigos únicos de medicamentos vigentes. La unidad de interés es una presentación CUM dentro de un expediente.",
        "Clasificar si una presentación corresponde a muestra médica usando atributos estructurales disponibles antes de conocer la etiqueta oficial. El resultado prioriza revisión académica; no sustituye al INVIMA.",
        "General: construir y desplegar una solución CRISP-DM reproducible. Específicos: auditar calidad y grano; controlar fuga; comparar cinco modelos y tres ensambles; ajustar hiperparámetros; evaluar en holdout y desplegar el pipeline.",
        "Se consolida a una fila por presentación, se divide 70/30 aislando expedientes, se valida con cinco folds agrupados, se balancea solo dentro del pipeline y se publica una app Streamlit que no reentrena.",
    ]
    for index, value in enumerate(business):
        set_cell(document.tables[1].cell(index, 1), value)
        shade(document.tables[1].cell(index, 0), LIGHT)
    set_cell(document.tables[2].cell(1, 0), "Predictivo")
    set_cell(document.tables[2].cell(1, 1), "Supervisado")
    set_cell(
        document.tables[2].cell(1, 2),
        "Logística, SVM, MLP, árbol, KNN; voting, bagging y AdaBoost",
    )
    set_cell(document.tables[2].cell(1, 3), "F1 clase Sí; apoyo: Accuracy, Precision, Recall, Balanced Accuracy, ROC-AUC y PR-AUC")

    dictionary_rows = [
        [
            row.variable,
            row.descripcion,
            f"{row.tipo_semantico} ({row.tipo_pandas_observado})",
            row.uso_modelado_y_tratamiento,
        ]
        for row in dictionary.itertuples(index=False)
    ]
    replace_table(
        document.tables[3],
        ["Variable", "Descripción", "Tipo", "Entrada/salida o tratamiento"],
        dictionary_rows,
    )

    numeric_rows = [
        [name, ">= 0; nulos imputados por mediana aprendida solo en train"]
        for name in [
            "cantidad_cum", "numero_filas_fuente",
            "numero_cantidades_ingrediente_distintas", "numero_principios_activos",
            "numero_vias_administracion", "numero_codigos_atc",
            "numero_unidades_medida", "numero_unidades_referencia",
            "numero_tipos_rol", "numero_actores",
        ]
    ]
    replace_table(document.tables[4], ["Variable", "Rango/regla esperada"], numeric_rows)
    categorical_rows = [
        [name, "Cadena no vacía; SIN_DATO si falta; categorías nuevas toleradas"]
        for name in [
            "forma_farmaceutica", "codigo_concentracion", "unidad_cum", "modalidad",
            "vias_administracion", "unidades_medida", "familias_atc", "tipos_rol",
        ]
    ]
    replace_table(document.tables[5], ["Variable", "Categorías válidas/regla"], categorical_rows)
    selected_rows = [
        [name, "Atributo estructural disponible al formular la presentación; sin información posterior ni identificadores"]
        for name, _ in numeric_rows + categorical_rows
    ]
    replace_table(document.tables[6], ["Variable seleccionada", "Justificación"], selected_rows)

    data = results["data"]
    statistical_rows = [
        ["Filas originales", fmt(data["raw_rows"])],
        ["Variables originales", fmt(data["raw_columns"])],
        ["Presentaciones CUM analíticas", fmt(data["modeling_rows"])],
        ["Expedientes únicos", fmt(data["unique_expedientes"])],
        ["Predictores", fmt(data["modeling_features"])],
        ["Variables numéricas", "10"],
        ["Variables categóricas", "8"],
        ["Clase positiva Sí", f"{fmt(data['target_positive_count'])} ({percentage(data['target_positive_rate'])})"],
        ["Clase negativa No", fmt(data["target_negative_count"])],
        ["Filas relacionales consolidadas", fmt(data["duplicate_expansion_rows"])],
        ["Reporte ydata", "reports/reporte_ydata_profiling.html — 155.807 filas completas"],
    ]
    replace_table(document.tables[7], ["Métrica", "Resultado"], statistical_rows)

    quality_rows = [
        ["Expansión ingrediente–rol–vía", "Consolidar por expedientecum + consecutivocum; no eliminar con drop_duplicates ciego"],
        ["Texto que contiene la respuesta", "Excluir producto y descripcioncomercial; la regla textual alcanzó 97,42% de accuracy artificial"],
        ["Fechas centinela 1000/1900/3000/9999", "Auditar y excluir del modelo por temporalidad/posible información posterior"],
        ["Variantes de unidades", "Normalizar tokens; no promediar cantidades expresadas en dimensiones incompatibles"],
        ["Caracteres de reemplazo/mojibake", "Cuantificar y conservar trazabilidad; no afecta los predictores estructurales seleccionados"],
        ["Objetivo", "Aceptar únicamente Si/No y detener ejecución ante etiquetas inesperadas o contradicción dentro del CUM"],
        ["Valores negativos", "Convertir a nulo y tratar dentro del pipeline; en el corte no se usan magnitudes de ingrediente mezcladas"],
        ["Categorías múltiples", "Separar por `|` y codificar multi-hot por token"],
    ]
    replace_table(document.tables[8], ["Tipo de error/riesgo", "Acción realizada"], quality_rows)

    null_rows = [
        [name, "0% o derivado del agregado", "Mediana dentro del pipeline; no se imputa objetivo/clave"]
        for name, _ in numeric_rows
    ] + [
        [name, "0% tras normalización a SIN_DATO", "Moda para simples; SIN_DATO y multi-hot para multivaluadas"]
        for name, _ in categorical_rows
    ]
    replace_table(document.tables[9], ["Variable", "% nulos analítico", "Método aplicado"], null_rows)
    encoding_rows = [
        ["es_muestra_medica", "Binaria: No=0, Sí=1; validación estricta"],
        ["Numéricas", "Mediana + log1p + StandardScaler, todo aprendido en cada fold"],
        ["Categóricas simples", "Moda + OneHotEncoder(handle_unknown='ignore', min_frequency=10)"],
        ["Categóricas multivaluadas", "CountVectorizer binario por token con min_df=10; no trata combinaciones como categorías atómicas"],
    ]
    replace_table(document.tables[10], ["Variable/grupo", "Método de codificación"], encoding_rows)
    reduction_rows = [
        ["Global", "No se aplica PCA/SVD para conservar interpretación y la estructura dispersa"],
        ["KNN", "TruncatedSVD de 30 componentes dentro de cada fold para distancias factibles en el espacio one-hot"],
    ]
    replace_table(document.tables[11], ["Técnica", "Justificación"], reduction_rows)
    replace_table(
        document.tables[12],
        ["Técnica aplicada", "Justificación"],
        [["RandomOverSampler", "Solo dentro de cada fold del 70% de entrenamiento; conserva intactos validación y test"]],
    )

    classic_names = [
        "Regresion logistica", "SVM lineal", "Red neuronal MLP",
        "Arbol de decision", "K vecinos mas cercanos",
    ]
    initial_params = {
        "Regresion logistica": "C=1; L2; liblinear",
        "SVM lineal": "C=1; squared_hinge; tol=0,001",
        "Red neuronal MLP": "(32); relu; alpha=0,001; early stopping",
        "Arbol de decision": "max_depth=12; min_samples_leaf=10",
        "K vecinos mas cercanos": "k=21; distance; SVD=30",
    }
    classic_rows = [
        [
            name,
            initial_params[name],
            "F1",
            fmt(results["cv_summary"][name]["f1"]["mean"]),
            fmt(results["cv_summary"][name]["f1"]["std"]),
        ]
        for name in classic_names
    ]
    replace_table(
        document.tables[13],
        ["Modelo", "Hiperparámetros iniciales", "Métrica", "CV Mean", "CV Std"],
        classic_rows,
    )
    best_classic = max(results["cv_summary"][name]["f1"]["mean"] for name in classic_names)
    ensemble_rows = []
    for name, kind in [
        ("Votacion", "Voting suave"),
        ("Bagging", "Bagging con árboles"),
        ("Boosting AdaBoost", "Boosting secuencial"),
    ]:
        score = results["cv_summary"][name]["f1"]["mean"]
        ensemble_rows.append([name, kind, fmt(score), fmt(score - best_classic)])
    replace_table(
        document.tables[14],
        ["Modelo", "Tipo", "F1 CV", "Mejora vs mejor clásico"],
        ensemble_rows,
    )

    tuning = results["hyperparameter_tuning"]
    logistic_params = tuning["Regresion logistica"]["best_params"]
    svm_params = tuning["SVM lineal"]["best_params"]
    tuning_rows = [
        ["Logística", "C", "Fuerza inversa de regularización", "0,1; 1; 10", str(logistic_params["model__C"])],
        ["Logística", "penalty (API: l1_ratio)", "L1 selecciona; L2 contrae", "L1/L2 (ratio 1/0)", tuning["Regresion logistica"]["best_penalty_alias"].upper()],
        ["SVM", "C", "Equilibrio margen/errores", "0,1; 1; 10", str(svm_params["model__C"])],
        ["SVM", "loss", "Penalización de violaciones", "hinge; squared_hinge", str(svm_params["model__loss"])],
    ]
    replace_table(
        document.tables[15],
        ["Modelo", "Hiperparámetro", "Qué controla", "Rango probado", "Valor final"],
        tuning_rows,
    )

    test_rows = [
        [
            row.modelo,
            fmt(row.accuracy), fmt(row.precision), fmt(row.recall), fmt(row.f1),
        ]
        for row in test.itertuples(index=False)
    ]
    replace_table(document.tables[16], ["Modelo", "Accuracy", "Precision", "Recall", "F1"], test_rows)
    replace_table(
        document.tables[17],
        ["Modelo", "MAE", "RMSE", "R²"],
        [["No aplica: problema de clasificación", "—", "—", "—"]],
    )
    selected = results["model_selection"]
    replace_table(
        document.tables[18],
        ["Modelo final", "Justificación técnica y estadística"],
        [[selected["selected_model"], selected["reason"] + " La selección no usó el test."]],
    )
    prediction_rows = []
    for row in predictions.head(5).itertuples(index=False):
        key = f"{row.expedientecum}-{row.consecutivocum}"
        prediction = "Sí" if int(row.prediccion) == 1 else "No"
        reality = "Sí" if int(row.es_muestra_medica) == 1 else "No"
        score = getattr(row, "puntaje_modelo_no_calibrado", None)
        prediction_rows.append(
            [key, f"Real={reality}; pred={prediction}; score={fmt(score)}", "Holdout aislado por expediente; puntaje no calibrado"]
        )
    replace_table(document.tables[19], ["Registro de prueba", "Predicción", "Interpretación"], prediction_rows)
    deployment_values = ["Sí", "Sí — validación local y pública", f"Sí — {args.streamlit_url}", "Sí — captura insertada en anexos"]
    for index, value in enumerate(deployment_values):
        set_cell(document.tables[20].cell(index, 1), value)
        shade(document.tables[20].cell(index, 0), LIGHT)
    set_table_width(document.tables[20], 2)

    document.add_page_break()
    add_heading(document, "7. Evidencia metodológica complementaria", 1)
    document.add_paragraph(
        "El formato se amplía para documentar controles exigidos por la guía y evitar "
        "conclusiones optimistas. Todas las cifras siguientes provienen de las salidas "
        "reproducibles incluidas en reports/ y del notebook público."
    )
    add_heading(document, "7.1 Granularidad, calidad y riesgo de fuga", 2)
    leak = quality["riesgo_fuga_particion"]
    add_table(
        document,
        ["Control", "Evidencia", "Decisión"],
        [
            ["Grano", f"{fmt(quality['dimensiones']['filas_originales'])} filas → {fmt(quality['dimensiones']['presentaciones_cum_unicas'])} CUM", "Agregar antes de modelar"],
            ["Split por filas incorrecto", f"{fmt(leak['porcentaje_filas_test_cuyo_cum_aparece_en_train'], 2)}% de filas test filtrarían CUM", "Prohibido"],
            ["Split por CUM, no por expediente", f"{fmt(leak['porcentaje_presentaciones_test_con_expediente_en_train_si_split_por_cum'], 2)}% compartirían expediente", "Agrupar por expedientecum"],
            ["Fuga textual", "Accuracy artificial 0,9742", "Excluir producto y descripcioncomercial"],
            ["Fechas/estados", "Contienen información temporal o posterior", "Excluir del modelo"],
            ["Objetivo", "0 CUM con etiqueta contradictoria", "Validación estricta Si/No"],
        ],
    )
    add_figure(document, figures / "multiplicidad_filas_por_cum.png", "Figura 1. Multiplicidad relacional del archivo original.")
    add_figure(document, figures / "faltantes_por_variable.png", "Figura 2. Valores faltantes por variable original.")

    add_heading(document, "7.2 Distribución, partición y línea base", 2)
    holdout = results["holdout"]
    add_table(
        document,
        ["Partición", "Filas", "%", "Tasa positiva", "Expedientes"],
        [
            ["Entrenamiento", fmt(holdout["train_rows"]), percentage(holdout["train_share"]), percentage(holdout["train_positive_rate"]), fmt(holdout["train_groups"])],
            ["Test intacto", fmt(holdout["test_rows"]), percentage(holdout["test_share"]), percentage(holdout["test_positive_rate"]), fmt(holdout["test_groups"])],
        ],
    )
    document.add_paragraph(
        f"Intersección de expedientes: {holdout['group_overlap']}. Accuracy de mayoría: "
        f"{fmt(data['majority_accuracy_baseline'])}. PR-AUC aleatoria: {fmt(data['pr_auc_random_baseline'])}."
    )
    add_figure(document, figures / "distribucion_objetivo.png", "Figura 3. Distribución del objetivo a nivel CUM.")

    add_heading(document, "7.3 Relaciones lineales y no lineales", 2)
    top_associations = associations.sort_values("asociacion_absoluta", ascending=False).head(12)
    add_table(
        document,
        ["Variable", "Tipo", "Método", "Asociación", "Nota"],
        [
            [row.variable, row.tipo, row.metodo, fmt(row.asociacion), "Exploratoria; no causal"]
            for row in top_associations.itertuples(index=False)
        ],
    )
    document.add_paragraph(
        "Las asociaciones se calcularon a nivel presentación CUM. Las variables con fuga "
        "o alta cardinalidad permanecen excluidas aunque exhiban asociación alta."
    )

    add_heading(document, "7.4 Validación cruzada y sobreajuste", 2)
    cv_columns = ["modelo", "f1_mean", "f1_std", "f1_train_mean", "roc_auc_mean", "pr_auc_mean"]
    available_cv_columns = [column for column in cv_columns if column in cv.columns]
    add_table(
        document,
        ["Modelo", "F1 media", "F1 std", "F1 train", "ROC-AUC", "PR-AUC"],
        [
            [
                row["modelo"], fmt(row.get("f1_mean")), fmt(row.get("f1_std")),
                fmt(row.get("f1_train_mean")), fmt(row.get("roc_auc_mean")), fmt(row.get("pr_auc_mean")),
            ]
            for _, row in cv.iterrows()
        ],
    )
    add_figure(document, figures / "comparacion_modelos_cv.png", "Figura 4. F1 medio y variabilidad en CV agrupada.")

    add_heading(document, "7.5 Comparación estadística y selección", 2)
    comparison = results["statistical_comparison"]
    add_table(
        document,
        ["Prueba/evidencia", "Resultado", "Interpretación"],
        [
            ["Friedman", f"χ²={fmt(comparison['friedman_statistic'])}; p={fmt(comparison['friedman_p'])}", "Contraste global no paramétrico sobre folds alineados"],
            ["Margen práctico", fmt(comparison["practical_f1_margin"], 2), "Se exige mejora >2 puntos de F1 frente a logística ajustada"],
            ["Mejor media CV", f"{comparison['best_cv_model']} · {fmt(comparison['best_cv_f1'])}", "No se decide por test"],
            ["Limitación", comparison["inference_limitation"], "La evidencia estadística es auxiliar"],
        ],
    )
    paired_rows = comparison["paired_comparisons_vs_tuned_logistic"]
    add_table(
        document,
        ["Candidato vs logística ajustada", "Δ F1", "p crudo", "p Holm"],
        [[row["candidate"], fmt(row["mean_difference_vs_tuned_logistic"]), fmt(row["raw_p"]), fmt(row["holm_adjusted_p"])] for row in paired_rows],
    )

    add_heading(document, "7.6 Evaluación final en test", 2)
    add_table(
        document,
        ["Modelo", "Accuracy", "Precision", "Recall", "F1", "Bal. Acc.", "ROC-AUC", "PR-AUC"],
        [[row.modelo, fmt(row.accuracy), fmt(row.precision), fmt(row.recall), fmt(row.f1), fmt(row.balanced_accuracy), fmt(row.roc_auc), fmt(row.pr_auc)] for row in test.itertuples(index=False)],
    )
    final_metrics = results["final_test_metrics"]
    matrix = results["final_confusion_matrix"]
    document.add_paragraph(
        f"Modelo congelado: {selected['selected_model']}. En test: F1={fmt(final_metrics['f1'])}, "
        f"precision={fmt(final_metrics['precision'])}, recall={fmt(final_metrics['recall'])}, "
        f"accuracy={fmt(final_metrics['accuracy'])}. Matriz [TN,FP;FN,TP]={matrix}."
    )
    add_figure(document, figures / "matriz_confusion_test.png", "Figura 5. Matriz de confusión del modelo congelado en test.", 5.2)
    add_figure(document, figures / "curva_roc_test.png", "Figura 6. Curva ROC en test.", 5.2)
    add_figure(document, figures / "curva_precision_recall_test.png", "Figura 7. Curva Precision-Recall en test.", 5.2)
    add_figure(document, figures / "coeficientes_modelo_final.png", "Figura 8. Pesos principales del modelo logístico; asociaciones, no causalidad.", 6.0)

    add_heading(document, "7.7 Despliegue y evidencia", 2)
    add_table(
        document,
        ["Artefacto", "Ubicación/enlace", "Estado"],
        [
            ["Dataset snapshot", "data/codigo_unico_medicamentos_vigentes_20260921.csv", "Incluido; SHA-256 verificable"],
            ["Notebook Colab", args.colab_url, "Público"],
            ["GitHub", args.github_url, "Público"],
            ["Streamlit", args.streamlit_url, "Activo"],
            ["Pipeline", "models/pipeline_muestra_medica.joblib", "Recargado y probado"],
            ["Metadata", "models/metadata_modelo.json", "Validado"],
            ["Profiling", "reports/reporte_ydata_profiling.html", "Completo: 155.807 filas"],
        ],
    )
    if args.screenshot and args.screenshot.exists():
        add_figure(document, args.screenshot, "Figura 9. Aplicación Streamlit pública y funcional.", 6.4)

    add_heading(document, "8. Conclusiones y limitaciones", 1)
    conclusions = [
        f"Se construyó un clasificador reproducible a nivel de presentación CUM; el modelo seleccionado fue {selected['selected_model']} con F1 de test {fmt(final_metrics['f1'])}.",
        "El aislamiento por expediente y el balanceo dentro de cada fold controlan las principales fuentes de fuga del archivo relacional.",
        "Los falsos negativos omiten posibles casos a priorizar; los falsos positivos generan revisión innecesaria. F1 equilibra ambos costos, pero el umbral debe validarse para una capacidad operativa real.",
        "El puntaje de la aplicación no es una probabilidad calibrada: RandomOverSampler cambia la prevalencia durante el ajuste. No debe comunicarse como riesgo individual.",
        "La fuente representa medicamentos vigentes en una fecha de corte. Cambios de catálogo, categorías nuevas o prácticas administrativas requieren monitoreo de deriva y reentrenamiento.",
        "El proyecto es académico: no evalúa seguridad, eficacia ni legalidad y no reemplaza la clasificación oficial del INVIMA.",
    ]
    for conclusion in conclusions:
        document.add_paragraph(f"• {conclusion}")

    add_heading(document, "Anexo A. Registro de prompts", 1)
    prompt_rows = get_prompt_rows(root / "notebooks" / "Proyecto_Final_ML_CUM.ipynb")
    add_table(document, ["ID", "Propósito", "Verificación humana"], prompt_rows)

    add_heading(document, "Anexo B. Trazabilidad técnica", 1)
    add_table(
        document,
        ["Elemento", "Valor"],
        [
            ["Fecha de corte", quality["fuente"]["fecha_corte_archivo"]],
            ["SHA-256 dataset", quality["fuente"]["sha256"]],
            ["Semilla", "42"],
            ["Split", "70/30 aproximado, estratificado y agrupado por expedientecum"],
            ["CV", "StratifiedGroupKFold, 5 folds, solo entrenamiento"],
            ["Pipeline", "Imputación + transformación + multi-hot + RandomOverSampler + modelo"],
            ["Reajuste despliegue", "Sí, con las 65.335 presentaciones solo después de cerrar la evaluación"],
        ],
    )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(f"Documento final guardado: {args.output}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--colab-url", required=True)
    parser.add_argument("--github-url", required=True)
    parser.add_argument("--streamlit-url", required=True)
    parser.add_argument("--screenshot", type=Path)
    return parser.parse_args()


if __name__ == "__main__":
    build_document(parse_args())
